"""
Розсилка пропозицій з твоєї пошти: лист + макет у вкладенні.

  python send_offers.py --test tvuj@email.cz   # один приклад собі – глянь, як виглядає
  python send_offers.py                        # показати, кому піде (нічого не відправляє)
  python send_offers.py --send                 # відправити (не більше daily_limit на день)

Налаштування SMTP і ліміти – у config.json. Хто вже отримав лист, записано в sent.json,
тому повторно нікому не піде.
"""
import argparse
import datetime as dt
import json
import random
import smtplib
import ssl
import time
from email.message import EmailMessage
from email.utils import formataddr, formatdate, make_msgid
from pathlib import Path

from openpyxl import load_workbook

from offers import load_config

HERE = Path(__file__).resolve().parent
SENT = HERE / "sent.json"


def load_offers(path):
    wb = load_workbook(path)
    ws = wb["Пропозиції"]
    header = [c.value for c in ws[1]]
    rows = []
    for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        d = dict(zip(header, r))
        d["_row"] = i
        rows.append(d)
    return wb, ws, header, rows


def build(cfg, row, img_dir, to=None):
    msg = EmailMessage()
    msg["From"] = formataddr((cfg["sender_name"], cfg["smtp_user"]))
    msg["To"] = to or row["Email"]
    msg["Reply-To"] = cfg["sender_email"]
    msg["Subject"] = row["Тема листа"]
    msg["Date"] = formatdate(localtime=True)
    msg["Message-ID"] = make_msgid(domain=cfg["smtp_user"].split("@")[-1])
    msg.set_content(row["Текст листа"])
    img = img_dir / row["Картинка"]
    msg.add_attachment(img.read_bytes(), maintype="image", subtype="jpeg",
                       filename="navrh-webu.jpg")
    return msg


def smtp_connect(cfg):
    ctx = ssl.create_default_context()
    if int(cfg["smtp_port"]) == 465:
        s = smtplib.SMTP_SSL(cfg["smtp_host"], 465, context=ctx, timeout=60)
    else:
        s = smtplib.SMTP(cfg["smtp_host"], int(cfg["smtp_port"]), timeout=60)
        s.starttls(context=ctx)
    s.login(cfg["smtp_user"], cfg["smtp_password"])
    return s


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--offers", default=str(HERE / "offers" / "nabidky.xlsx"))
    ap.add_argument("--send", action="store_true", help="реально відправити")
    ap.add_argument("--test", metavar="EMAIL", help="відправити один приклад на свою адресу")
    ap.add_argument("--status", default="", help="тільки цей пріоритет, напр. 'БЕЗ САЙТУ'")
    args = ap.parse_args()

    cfg = load_config()
    offers = Path(args.offers)
    wb, ws, header, rows = load_offers(offers)
    sent = json.loads(SENT.read_text("utf8")) if SENT.exists() else {}
    today = dt.date.today().isoformat()

    if args.test:
        s = smtp_connect(cfg)
        s.send_message(build(cfg, rows[0], offers.parent, to=args.test))
        s.quit()
        print(f"Тестовий лист для «{rows[0]['Назва']}» відправлено на {args.test}")
        return

    queue, seen = [], set(sent)
    for r in rows:
        email = (r.get("Email") or "").strip().lower()
        if not email or email in seen:
            continue
        if args.status and r.get("Пріоритет") != args.status:
            continue
        seen.add(email)
        queue.append(r)

    left_today = cfg["daily_limit"] - sum(1 for d in sent.values() if d == today)
    batch = queue[:max(left_today, 0)]
    print(f"В черзі: {len(queue)} | сьогодні ще можна: {max(left_today, 0)} | відправлю: {len(batch)}")
    for r in batch:
        print(f"  → {r['Email']:35} {r['Назва']}")
    if not args.send:
        print("\nЦе був перегляд. Щоб відправити: python send_offers.py --send")
        return

    sent_col = header.index("Відправлено") + 1
    s = smtp_connect(cfg)
    for n, r in enumerate(batch, 1):
        try:
            s.send_message(build(cfg, r, offers.parent))
        except smtplib.SMTPServerDisconnected:
            s = smtp_connect(cfg)
            s.send_message(build(cfg, r, offers.parent))
        except smtplib.SMTPRecipientsRefused:
            print(f"  ✗ адресу відхилено: {r['Email']}")
            sent[r["Email"].strip().lower()] = "bad"
            continue
        sent[r["Email"].strip().lower()] = today
        SENT.write_text(json.dumps(sent, ensure_ascii=False, indent=1), "utf8")
        ws.cell(r["_row"], sent_col).value = today
        print(f"  ✓ {n}/{len(batch)} {r['Email']}")
        if n < len(batch):
            time.sleep(random.uniform(*cfg["delay_seconds"]))  # пауза, щоб не виглядати як спамер
    s.quit()
    try:
        wb.save(offers)
    except PermissionError:
        print("! Не зміг оновити nabidky.xlsx (відкритий в Excel?) – але sent.json записано.")
    print("Готово.")


if __name__ == "__main__":
    main()
