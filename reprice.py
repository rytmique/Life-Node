"""
Перерахувати ціни в уже готовому Excel (наприклад після зміни NICHE_PRICING у lead_finder.py).
Підтягує правову форму фірми з ARES за IČO.
Запуск:  python reprice.py output/leady_2026-10-18_1156.xlsx
"""
import sys
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from openpyxl import load_workbook

import lead_finder as lf

REVERSE = {"так": True, "ні": False, "?": None}


def load(path, sheet="Усі фірми"):
    ws = load_workbook(path)[sheet]
    header = [c.value for c in ws[1]]
    by_title = {t: k for t, k, _ in lf.COLUMNS}
    rows = []
    for r in ws.iter_rows(min_row=2):
        c = {}
        for title, cell in zip(header, r):
            key = by_title.get(title)
            if key:
                c[key] = cell.hyperlink.target if key in ("firmy_url", "image") and cell.hyperlink else cell.value
        c = {k: ("" if v is None else v) for k, v in c.items()}
        c["email_ok"] = REVERSE.get(c.get("email_ok"))
        c["new"] = c.get("new") == "NEW"
        c["ico"] = str(c.get("ico") or "").split(".")[0].zfill(8) if c.get("ico") else ""
        rows.append(c)
    return rows


def enrich(c):
    if not c.get("legal_form"):
        _, name, form = lf.ares_check(c["ico"])
        c["legal_form"] = form
        c["legal_name"] = c.get("legal_name") or name
    return lf.price(c)


def main():
    src = Path(sys.argv[1])
    rows = load(src)
    print(f"{len(rows)} фірм, тягну правову форму з ARES …")
    with ThreadPoolExecutor(8) as ex:
        rows = list(ex.map(enrich, rows))
    rows.sort(key=lambda l: (lf.PRIORITY[l["status"]], -int(l["score"] or 0),
                             l["phone_type"] != "мобільний", -int(l["reviews"] or 0)))
    out = src.with_name(src.stem + "_ceny.xlsx")
    lf.write_excel(rows, out)
    leads = [r for r in rows if r["status"] != "OK"]
    total = sum(r["price"] for r in leads)
    print(f"Готово: {out}\n  лідів: {len(leads)}, сума всіх цін: {total:,} Kč, "
          f"середня: {total // max(len(leads), 1):,} Kč")


if __name__ == "__main__":
    main()
