"""
Пошук клієнтів для веб-студії в Чехії.

Джерело: Firmy.cz (каталог Seznam) -> перевірка сайту -> перевірка в ARES -> Excel.
Запуск:  python lead_finder.py --cities Brno Praha --niches "kadeřnictví" "autoservis"
"""
import argparse
import datetime as dt
import html as htmllib
import json
import random
import re
import socket
import sys
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from urllib.parse import quote_plus, urljoin, urlparse

import requests
import urllib3
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    import dns.resolver  # pip install dnspython
except ImportError:
    dns = None

urllib3.disable_warnings()

# ---------------------------------------------------------------- налаштування
DEFAULT_CITIES = ["Praha", "Brno", "Ostrava", "Plzeň", "Liberec", "Olomouc",
                  "České Budějovice", "Hradec Králové", "Pardubice", "Zlín"]

DEFAULT_NICHES = [
    "kadeřnictví", "kosmetický salon", "manikúra", "masáže", "barbershop",
    "autoservis", "pneuservis", "autoškola", "zubní ordinace", "veterinář",
    "restaurace", "kavárna", "pekárna", "květinářství", "fotograf",
    "instalatér", "elektrikář", "truhlářství", "malíř pokojů", "stavební firma",
    "účetní", "fitness", "realitní kancelář", "úklidové služby", "zámečnictví",
]

UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/130.0 Safari/537.36")
FIRMY_DELAY = (0.8, 1.6)      # пауза між запитами до firmy.cz, сек
SITE_TIMEOUT = 15
THIS_YEAR = dt.date.today().year

SOCIAL = ("facebook.com", "fb.com", "instagram.com", "linkedin.com", "youtube.com",
          "tiktok.com", "twitter.com", "x.com", "linktr.ee", "google.com", "goo.gl",
          "mapy.cz", "mapy.com", "firmy.cz", "booking.com", "reservio.com",
          "reservio.cz", "bookio.com", "noona.app")
FREE_BUILDERS = ("webnode.", "estranky.", "wbs.cz", "mypage.cz", "wixsite.com",
                 "sites.google.com", "blogspot.", "websnadno.cz", "webmium.",
                 "mozello.", "jimdo", "weebly.com", "rajce.", "unas.cz",
                 "own.cz", "ic.cz", "wz.cz", "g6.cz", "4fan.cz", "tode.cz")
PARKED = ("doména je na prodej", "domena je na prodej", "this domain is for sale",
          "doména byla zaregistrována", "stránka ve výstavbě", "stranka ve vystavbe",
          "under construction", "web se připravuje", "připravujeme nový web",
          "parkovaná doména", "domain parking", "coming soon", "website expired")
JUNK_EMAIL = ("example.", "sentry", "wixpress", "domain.", "email.cz@", "@2x",
              ".png", ".jpg", ".gif", ".webp", ".svg", "yourname", "jmeno@",
              "vas@", "info@firmy.cz", "@seznam-firmy", "napoveda")

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
LD_RE = re.compile(r'<script type="application/ld\+json">(.*?)</script>', re.S)

_local = threading.local()
_firmy_lock = threading.Lock()
_firmy_last = [0.0]


def session():
    if not hasattr(_local, "s"):
        s = requests.Session()
        s.headers.update({"User-Agent": UA, "Accept-Language": "cs,en;q=0.8"})
        _local.s = s
    return _local.s


def get(url, tries=3, **kw):
    kw.setdefault("timeout", 40)
    for i in range(tries):
        try:
            r = session().get(url, **kw)
            if r.status_code in (429, 502, 503, 504) and i < tries - 1:
                time.sleep(5 * (i + 1))
                continue
            return r
        except requests.RequestException:
            if i == tries - 1:
                raise
            time.sleep(3 * (i + 1))


def firmy_get(url):
    """Ввічливий запит до firmy.cz: не частіше ніж раз на ~1 сек."""
    with _firmy_lock:
        wait = _firmy_last[0] + random.uniform(*FIRMY_DELAY) - time.time()
        if wait > 0:
            time.sleep(wait)
        _firmy_last[0] = time.time()
    return get(url)


# ---------------------------------------------------------------- firmy.cz
def search_firmy(query, max_pages):
    """Повертає список карток з видачі firmy.cz."""
    found, seen = [], set()
    for page in range(1, max_pages + 1):
        # firmy.cz редиректить "Brno"->"brno" і "page=1"->головна, губляючи пагінацію
        url = f"https://www.firmy.cz/?potrebuji={quote_plus(query.lower())}"
        if page > 1:
            url += f"&page={page}"
        try:
            r = firmy_get(url)
        except requests.RequestException as e:
            print(f"   ! {query} p{page}: {e}")
            break
        if r.status_code != 200:
            break
        new = 0
        for block in r.text.split('<article class="premiseBox')[1:]:
            m = re.search(r'id="prem-(\d+)"', block)
            ld = LD_RE.search(block)
            if not m or not ld or m.group(1) in seen:
                continue
            try:
                data = json.loads(ld.group(1))
            except json.JSONDecodeError:
                continue
            seen.add(m.group(1))
            new += 1
            found.append(card_from_ld(m.group(1), data))
        if new == 0:
            break
    return found


def card_from_ld(pid, d):
    addr = d.get("address") or {}
    web = ""
    for u in d.get("sameAs") or []:
        if "firmy.cz" not in u:
            web = u
            break
    rating = d.get("aggregateRating") or {}
    return {
        "id": pid,
        "name": htmllib.unescape(d.get("name", "")).strip(),
        "phone": d.get("telephone", ""),
        "street": addr.get("streetAddress", ""),
        "city": addr.get("addressLocality", ""),
        "zip": addr.get("postalCode", ""),
        "website": web,
        "firmy_url": d.get("url", ""),
        "reviews": rating.get("ratingCount", 0) or 0,
        "rating": round(int(rating.get("ratingValue", 0) or 0) / 20, 1) if rating else "",
        "description": (d.get("description") or "")[:300],
        "image": d.get("image", ""),
    }


def enrich_from_detail(c):
    """Email, IČO і точніший сайт зі сторінки фірми."""
    try:
        r = firmy_get(c["firmy_url"])
    except requests.RequestException:
        return
    h = r.text
    emails = [e for e in re.findall(r'mailto:([^"?]+)', h) if good_email(e)]
    if emails:
        c["email"] = htmllib.unescape(emails[0]).lower()
    m = re.search(r'detailBusinessInfo ico">\s*(\d{6,8})', h)
    if m:
        c["ico"] = m.group(1).zfill(8)
    for ld in LD_RE.findall(h):
        try:
            d = json.loads(ld)
        except json.JSONDecodeError:
            continue
        if d.get("@type") == "LocalBusiness":
            u = d.get("url", "")
            if u and "firmy.cz" not in u and not c["website"]:
                c["website"] = u
            if d.get("telephone"):
                c["phone"] = d["telephone"]


# ---------------------------------------------------------------- контакти
def good_email(e):
    e = e.lower().strip()
    return bool(EMAIL_RE.fullmatch(e)) and not any(j in e for j in JUNK_EMAIL)


_mx_cache = {}
_mx_works = None


def mx_available():
    """Чи працює DNS-перевірка взагалі (у деяких мережах MX-запити блокують)."""
    global _mx_works
    if _mx_works is None:
        try:
            _mx_works = dns is not None and bool(dns.resolver.resolve("seznam.cz", "MX", lifetime=8))
        except Exception:
            _mx_works = False
    return _mx_works


def email_has_mx(email):
    """True = домен приймає пошту, False = домену/пошти немає, None = не вдалось перевірити."""
    if not email or not mx_available():
        return None
    domain = email.split("@")[-1]
    if domain not in _mx_cache:
        try:
            dns.resolver.resolve(domain, "MX", lifetime=8)
            _mx_cache[domain] = True
        except dns.resolver.NXDOMAIN:
            _mx_cache[domain] = False
        except dns.resolver.NoAnswer:
            _mx_cache[domain] = domain_exists(domain)  # без MX пошта йде на A-запис
        except Exception:
            _mx_cache[domain] = None
    return _mx_cache[domain]


def domain_exists(host):
    try:
        socket.getaddrinfo(host, 443)
        return True
    except socket.gaierror:
        return False


def normalize_phone(p):
    digits = re.sub(r"\D", "", p or "")
    if digits.startswith("00420"):
        digits = digits[5:]
    elif digits.startswith("420") and len(digits) == 12:
        digits = digits[3:]
    if len(digits) != 9:
        return (p or "").strip(), ""
    kind = "мобільний" if digits[0] in "67" else "стаціонарний"
    return f"+420 {digits[:3]} {digits[3:6]} {digits[6:]}", kind


def ares_check(ico):
    """(чи фірма ще існує, офіційна назва (для ОСВЧ – ім'я власника), код правової форми)."""
    if not ico:
        return None, "", ""
    try:
        r = get(f"https://ares.gov.cz/ekonomicke-subjekty-v-be/rest/ekonomicke-subjekty/{ico}",
                timeout=20)
        if r.status_code == 404:
            return False, "", ""
        d = r.json()
        return d.get("datumZaniku") is None, d.get("obchodniJmeno", ""), str(d.get("pravniForma") or "")
    except Exception:
        return None, "", ""


# ---------------------------------------------------------------- аудит сайту
def audit_site(url):
    """Повертає (статус, бали, причини[], emails[])."""
    if not url:
        return "БЕЗ САЙТУ", 100, ["немає сайту"], []
    if not url.startswith("http"):
        url = "http://" + url
    host = urlparse(url).netloc.lower()
    if any(s in host for s in SOCIAL):
        return "БЕЗ САЙТУ", 95, [f"замість сайту лише {host}"], []

    reasons, score = [], 0
    t0 = time.time()
    try:
        r = get(url, tries=2, timeout=SITE_TIMEOUT, allow_redirects=True)
    except requests.exceptions.SSLError:
        try:
            r = get(url, tries=1, timeout=SITE_TIMEOUT, verify=False)
            reasons.append("битий SSL-сертифікат (браузер лякає клієнтів)")
            score += 30
        except requests.RequestException:
            return site_down(host)
    except requests.RequestException:
        return site_down(host)
    elapsed = time.time() - t0

    if r.status_code in (401, 403, 429):  # антибот-захист, сайт швидше за все живий
        return "OK", 0, [f"не вдалось перевірити (HTTP {r.status_code}), глянь вручну"], []
    if r.status_code >= 400:
        return "САЙТ НЕ ПРАЦЮЄ", 90, [f"сайт видає помилку {r.status_code}"], []

    h = r.text or ""
    low = h.lower()
    final = r.url.lower()
    text = re.sub(r"<script.*?</script>|<style.*?</style>|<[^>]+>", " ", low, flags=re.S)
    text = re.sub(r"\s+", " ", text)

    if any(p in text for p in PARKED) and len(text) < 3000:
        return "БЕЗ САЙТУ", 90, ["на домені заглушка / сайт у розробці"], []
    if any(b in final for b in FREE_BUILDERS):
        reasons.append("безкоштовний конструктор з чужим доменом")
        score += 30
    if not final.startswith("https"):
        reasons.append("немає HTTPS ('Nezabezpečeno' у браузері)")
        score += 20
    if 'name="viewport"' not in low and "name='viewport'" not in low and "name=viewport" not in low:
        reasons.append("не адаптований під телефон")
        score += 30

    years = [int(y) for y in re.findall(
        r"(?:©|&copy;|copyright|\(c\))\s*(?:[^<]{0,40}?)((?:19|20)\d\d)(?!\d)", low)]
    years = [y for y in years if 1995 <= y <= THIS_YEAR]
    if years and max(years) <= THIS_YEAR - 4:
        reasons.append(f"© {max(years)} – сайт роками не оновлювали")
        score += 15 if max(years) > THIS_YEAR - 7 else 25

    m = re.search(r"jquery[.-]?(\d)\.(\d+)", low)
    if m and (int(m.group(1)), int(m.group(2))) < (1, 9):
        reasons.append(f"застарілий jQuery {m.group(1)}.{m.group(2)}")
        score += 10
    gen = re.search(r'<meta[^>]+name=["\']generator["\'][^>]+content=["\']([^"\']+)', low)
    if gen:
        g = gen.group(1)
        old = (re.search(r"wordpress ([1-4])\.", g) or re.search(r"joomla! ?(1|2|3)\.", g)
               or "frontpage" in g or "dreamweaver" in g or g.startswith("microsoft word"))
        if old:
            reasons.append(f"стара CMS ({g[:30]})")
            score += 15
    if '.swf"' in low or "<marquee" in low or "<frameset" in low:
        reasons.append("технології 2000-х (flash/frames)")
        score += 20
    if low.count("<table") >= 3 and "<div" not in low[:5000]:
        reasons.append("верстка таблицями")
        score += 10
    if elapsed > 5:
        reasons.append(f"повільний ({elapsed:.0f} с)")
        score += 10
    spa = any(k in low for k in ('id="root"', 'id="app"', "__next", "nuxt"))
    if len(text) < 400 and not spa:
        reasons.append("майже порожній сайт")
        score += 15
    if "<title" not in low or re.search(r"<title>\s*</title>", low):
        reasons.append("немає заголовка сторінки (погано для Google)")
        score += 5

    emails = collect_emails(h, r.url)
    score = min(score, 89)
    status = "СТАРИЙ САЙТ" if score >= 35 else "СЛАБКИЙ САЙТ" if score >= 20 else "OK"
    return status, score, reasons, emails


def site_down(host):
    if not domain_exists(host):
        return "САЙТ НЕ ПРАЦЮЄ", 92, ["домен не існує / прострочений"], []
    return "САЙТ НЕ ПРАЦЮЄ", 85, ["сайт не відкривається (тайм-аут)"], []


def collect_emails(h, base):
    emails = re.findall(r'mailto:([^"\'?>\s]+)', h) + EMAIL_RE.findall(h)
    emails = [htmllib.unescape(e).lower() for e in emails if good_email(e)]
    if not emails:
        for path in ("/kontakt", "/kontakty", "/contact"):
            try:
                r = get(urljoin(base, path), tries=1, timeout=10, verify=False)
            except requests.RequestException:
                continue
            if r.status_code == 200:
                emails = [e.lower() for e in EMAIL_RE.findall(r.text) if good_email(e)]
                if emails:
                    break
    return list(dict.fromkeys(emails))[:3]


# ---------------------------------------------------------------- pitch
def pitch(status, reasons):
    if status == "БЕЗ САЙТУ":
        if reasons and "лише" in reasons[0]:
            return ("Dobrý den, našel jsem vás na Firmy.cz – máte jen profil na sociální síti. "
                    "Vlastní web vám přivede zákazníky z Googlu. Připravím vám návrh zdarma.")
        return ("Dobrý den, našel jsem vás na Firmy.cz, ale nemáte webové stránky. "
                "Zákazníci vás tak na Googlu nenajdou. Rád vám připravím návrh webu zdarma.")
    if status == "САЙТ НЕ ПРАЦЮЄ":
        return ("Dobrý den, zkoušel jsem otevřít váš web a nefunguje – zákazníci vidí chybu. "
                "Můžu vám pomoct s novým funkčním webem, návrh zdarma.")
    r = " ".join(reasons)
    if "телефон" in r:
        return ("Dobrý den, díval jsem se na váš web – na mobilu se špatně zobrazuje, "
                "a přitom většina zákazníků hledá z telefonu. Připravím vám moderní návrh zdarma.")
    if "HTTPS" in r or "SSL" in r:
        return ("Dobrý den, prohlížeč u vašeho webu ukazuje 'Nezabezpečeno', což odrazuje zákazníky. "
                "Rád vám udělám nový bezpečný a moderní web.")
    return ("Dobrý den, váš web působí zastarale a nějakou dobu nebyl aktualizován. "
            "Připravím vám zdarma návrh moderního webu.")


# ---------------------------------------------------------------- ціна
# Два пакети для всіх. Правиш тут.
START_PRICE = 4990      # односторінковий сайт: послуги, контакти, мапа, форма
STANDARD_PRICE = 7990   # повний сайт як на макеті: + запис / меню / галерея
STANDARD_EXTRAS = {  # що додає Standard для ніші
    "booking": ("сайт + онлайн-запис + ціни", (
        "kadeřnictví", "kosmetický salon", "manikúra", "masáže", "barbershop", "fitness", "autoškola",
        "zubní ordinace", "veterinář", "pneuservis")),
    "food": ("сайт + меню + бронювання столу + галерея", ("restaurace", "kavárna", "pekárna")),
}
DEFAULT_EXTRAS = "сайт + послуги + галерея робіт + форма запиту"
STATUS_WORK = {"БЕЗ САЙТУ": "новий сайт + домен", "САЙТ НЕ ПРАЦЮЄ": "новий сайт + відновлення домену",
               "СТАРИЙ САЙТ": "редизайн + перенос контенту", "СЛАБКИЙ САЙТ": "модернізація сайту",
               "OK": "—"}
# коди правової форми ARES: 101-108 ОСВЧ, 112 s.r.o., 121 a.s.
LEGAL_K = {"112": (1.15, "s.r.o."), "121": (1.4, "a.s."), "113": (1.1, "k.s."),
           "111": (1.1, "v.o.s.")}


def legal_form(code, name):
    if code in LEGAL_K:
        return LEGAL_K[code]
    if code.startswith("10"):
        return 1.0, "ОСВЧ"
    n = (name or "").lower()
    if "a.s." in n:
        return LEGAL_K["121"]
    if "s.r.o" in n or "spol. s r" in n:
        return LEGAL_K["112"]
    return (1.05, "інше") if code else (1.0, "?")


def price(c):
    """Ціни пакетів Start / Standard + підтримка за розміром фірми."""
    extras = DEFAULT_EXTRAS
    for pkg, niches in STANDARD_EXTRAS.values():
        if c.get("niche") in niches:
            extras = pkg
            break
    k_legal, form = legal_form(c.get("legal_form", ""), c.get("legal_name") or c.get("name"))
    c["price_start"] = START_PRICE
    c["price"] = STANDARD_PRICE
    c["support"] = 490 if k_legal >= 1.4 else 390 if k_legal > 1 else 290
    c["package"] = f"{STATUS_WORK.get(c['status'], '')}: {extras}"
    c["form"] = form
    return c


# ---------------------------------------------------------------- обробка
def process(c):
    enrich_from_detail(c)
    status, score, reasons, site_emails = audit_site(c["website"])
    c["status"], c["score"], c["reasons"] = status, score, "; ".join(reasons)
    if not c.get("email") and site_emails:
        c["email"] = site_emails[0]
    c["email"] = c.get("email", "")
    c["email_ok"] = email_has_mx(c["email"])
    c["phone"], c["phone_type"] = normalize_phone(c["phone"])
    c["active"], c["legal_name"], c["legal_form"] = ares_check(c.get("ico", ""))
    c["pitch"] = pitch(status, reasons)
    return price(c)


PRIORITY = {"БЕЗ САЙТУ": 1, "САЙТ НЕ ПРАЦЮЄ": 2, "СТАРИЙ САЙТ": 3, "СЛАБКИЙ САЙТ": 4, "OK": 9}
FILLS = {"БЕЗ САЙТУ": "FFC7CE", "САЙТ НЕ ПРАЦЮЄ": "FFD8A8", "СТАРИЙ САЙТ": "FFEB9C",
         "СЛАБКИЙ САЙТ": "E2EFDA", "OK": "FFFFFF"}

COLUMNS = [  # (заголовок, ключ, ширина)
    ("Пріоритет", "status", 16), ("Бали", "score", 6), ("Новий", "new", 7),
    ("Назва", "name", 34), ("Start, Kč", "price_start", 10), ("Standard, Kč", "price", 11),
    ("Підтримка/міс, Kč", "support", 10), ("Що в Standard", "package", 40), ("Форма", "form", 8), ("Телефон", "phone", 17), ("Тип тел.", "phone_type", 13),
    ("Email", "email", 30), ("Email живий", "email_ok", 11), ("Сайт", "website", 30),
    ("Що не так", "reasons", 45), ("Ніша", "niche", 16), ("Місто", "city", 20),
    ("Адреса", "street", 26), ("PSČ", "zip", 7), ("Відгуки", "reviews", 8),
    ("Рейтинг", "rating", 8), ("IČO", "ico", 10), ("Власник / юр. назва (ARES)", "legal_name", 28),
    ("Firmy.cz", "firmy_url", 14), ("Текст для першого контакту (CZ)", "pitch", 60),
    ("Статус (твій)", "my_status", 14), ("Нотатки", "notes", 25),
    ("Опис (Firmy.cz)", "description", 30), ("Фото (Firmy.cz)", "image", 12),
]


def yes_no(v):
    return "так" if v is True else "ні" if v is False else "?"


def write_excel(leads, path):
    wb = Workbook()
    keys = [c[1] for c in COLUMNS]
    wrap_cols = {keys.index(k) + 1 for k in ("reasons", "pitch", "package")}
    money_cols = {keys.index(k) + 1 for k in ("price_start", "price", "support")}
    sheets = [("Ліди", [l for l in leads if l["status"] != "OK"]), ("Усі фірми", leads)]
    for i, (title, rows) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        ws.append([c[0] for c in COLUMNS])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="305496")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for l in rows:
            vals = []
            for _, key, _ in COLUMNS:
                v = l.get(key, "")
                if key in ("email_ok", "new"):
                    v = yes_no(v) if key == "email_ok" else ("NEW" if v else "")
                vals.append(v)
            ws.append(vals)
            row = ws.max_row
            fill = PatternFill("solid", fgColor=FILLS[l["status"]])
            for col in range(1, len(COLUMNS) + 1):
                ws.cell(row, col).fill = fill
                ws.cell(row, col).alignment = Alignment(vertical="top", wrap_text=col in wrap_cols)
                if col in money_cols:
                    ws.cell(row, col).number_format = "#,##0"
            for key in ("website", "firmy_url", "image"):
                col = [c[1] for c in COLUMNS].index(key) + 1
                u = l.get(key)
                if u:
                    cell = ws.cell(row, col)
                    cell.hyperlink = u if u.startswith("http") else "http://" + u
                    cell.font = Font(color="0563C1", underline="single")
                    if key != "website":
                        cell.value = "відкрити"
            if l.get("phone"):
                col = [c[1] for c in COLUMNS].index("phone") + 1
                ws.cell(row, col).hyperlink = "tel:" + l["phone"].replace(" ", "")
        for idx, (_, _, w) in enumerate(COLUMNS, 1):
            ws.column_dimensions[get_column_letter(idx)].width = w
        ws.freeze_panes = "E2"
        ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description="Пошук клієнтів без сайту / зі старим сайтом (CZ)")
    ap.add_argument("--cities", nargs="+", default=DEFAULT_CITIES)
    ap.add_argument("--niches", nargs="+", default=DEFAULT_NICHES)
    ap.add_argument("--max-pages", type=int, default=5, help="сторінок видачі на запит (по ~15 фірм)")
    ap.add_argument("--workers", type=int, default=6)
    ap.add_argument("--out", default="")
    ap.add_argument("--include-inactive", action="store_true", help="не викидати закриті за ARES фірми")
    args = ap.parse_args()

    here = Path(__file__).resolve().parent
    out_dir = here / "output"
    out_dir.mkdir(exist_ok=True)
    db_path = here / "seen.json"
    seen = json.loads(db_path.read_text("utf8")) if db_path.exists() else {}

    cards = {}
    for city in args.cities:
        for niche in args.niches:
            q = f"{niche} {city}"
            res = search_firmy(q, args.max_pages)
            for c in res:
                c["niche"] = niche
                cards.setdefault(c["id"], c)
            print(f"[пошук] {q}: {len(res)} фірм (усього унікальних: {len(cards)})")

    print(f"[аналіз] перевіряю {len(cards)} фірм: сайт, email, ARES …")
    leads, done = [], 0
    with ThreadPoolExecutor(args.workers) as ex:
        futs = [ex.submit(process, c) for c in cards.values()]
        for f in as_completed(futs):
            done += 1
            try:
                c = f.result()
            except Exception as e:
                print(f"   ! помилка: {e}")
                continue
            if c["active"] is False and not args.include_inactive:
                continue
            if not c["phone"] and not c["email"]:
                continue  # без контактів нам ні до чого
            leads.append(c)
            if done % 25 == 0:
                print(f"   {done}/{len(futs)}")

    uniq, keys = [], set()
    for l in sorted(leads, key=lambda l: -(l["reviews"] or 0)):
        k = urlparse(l["website"]).netloc.replace("www.", "") if l["website"] else l["phone"]
        if k and k in keys and not any(s in k for s in SOCIAL):
            continue
        keys.add(k)
        uniq.append(l)
    leads = uniq

    today = dt.date.today().isoformat()
    for l in leads:
        l["new"] = l["id"] not in seen
        seen.setdefault(l["id"], today)
    db_path.write_text(json.dumps(seen, ensure_ascii=False), "utf8")

    leads.sort(key=lambda l: (PRIORITY[l["status"]], -l["score"], l["phone_type"] != "мобільний",
                              -(l["reviews"] or 0)))
    out = Path(args.out) if args.out else out_dir / f"leady_{today}_{dt.datetime.now():%H%M}.xlsx"
    write_excel(leads, out)

    stats = {}
    for l in leads:
        stats[l["status"]] = stats.get(l["status"], 0) + 1
    print("\nГотово:", out)
    for s in sorted(stats, key=PRIORITY.get):
        print(f"  {s:15} {stats[s]}")
    print(f"  нових з минулого запуску: {sum(l['new'] for l in leads)}")


if __name__ == "__main__":
    sys.exit(main())
