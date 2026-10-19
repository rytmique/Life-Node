"""
Персональна пропозиція для кожного ліда: унікальний макет сайту 1920x1080 (JPG) + лист чеською.

Запуск:  python offers.py output/leady_2026-10-18_1156_ceny.xlsx
         python offers.py output/leady_….xlsx --limit 20      # тільки перші 20
Результат: папка offers/ з картинками + offers/nabidky.xlsx (тема, текст листа, WhatsApp, файл).
Свої дані (ім'я, телефон) – у config.json.
"""
import argparse
import html
import json
import os
import re
import unicodedata
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import lead_finder as lf
from reprice import load

HERE = Path(__file__).resolve().parent
CONFIG = HERE / "config.json"
DEFAULT_CONFIG = {
    "sender_name": "Tvoje Jméno",
    "sender_phone": "+420 000 000 000",
    "sender_email": "tvuj@email.cz",
    "sender_web": "",
    "delivery_days": 7,
    "smtp_host": "smtp.gmail.com",
    "smtp_port": 587,
    "smtp_user": "tvuj@email.cz",
    "smtp_password": "APP-PASSWORD",
    "daily_limit": 25,
    "delay_seconds": [90, 240],
}

# ---------------------------------------------------------------- вміст по нішах
# tagline, CTA, 4 послуги (назва, опис), палітра, віджет у hero
NICHES = {
    "kadeřnictví": ("Účes, se kterým se budete cítit skvěle", "Objednat se online",
                    [("Dámské střihy", "Střih na míru tvaru obličeje"), ("Barvení & melír", "Balayage, ombré, tónování"),
                     ("Pánské střihy", "Klasika i moderní styl"), ("Společenské účesy", "Svatby, plesy, focení")],
                    "beauty", "booking"),
    "kosmetický salon": ("Péče o pleť s viditelným výsledkem", "Rezervovat termín",
                         [("Ošetření pleti", "Čištění, hydratace, anti-age"), ("Lifting řas", "Přirozeně zvýrazněný pohled"),
                          ("Úprava obočí", "Tvarování a barvení"), ("Masáž obličeje", "Relax a regenerace")],
                         "beauty", "booking"),
    "manikúra": ("Krásné nehty bez čekání", "Rezervovat termín",
                 [("Gelové nehty", "Modeláž a doplnění"), ("Manikúra", "Klasická i japonská"),
                  ("Pedikúra", "Mokrá i suchá"), ("Nail art", "Zdobení dle přání")], "beauty", "booking"),
    "masáže": ("Odpočinek pro tělo i mysl", "Rezervovat masáž",
               [("Klasická masáž", "Uvolnění zad a šíje"), ("Relaxační masáž", "Aromaterapie a klid"),
                ("Sportovní masáž", "Regenerace po výkonu"), ("Dárkové poukazy", "Radost pro vaše blízké")],
               "spa", "booking"),
    "barbershop": ("Střih, vousy a styl", "Objednat se",
                   [("Pánský střih", "Fade, klasika, nůžky"), ("Úprava vousů", "Tvarování a péče"),
                    ("Hot towel", "Holení břitvou"), ("Dětský střih", "Pro malé gentlemany")], "dark", "booking"),
    "autoservis": ("Spolehlivý servis vašeho auta", "Objednat servis",
                   [("Servisní prohlídky", "Dle předpisu výrobce"), ("Diagnostika", "Rychlé odhalení závad"),
                    ("Brzdy a podvozek", "Opravy a výměny"), ("Příprava na STK", "Bez zbytečných starostí")],
                   "auto", "quote"),
    "pneuservis": ("Přezutí bez front a čekání", "Rezervovat přezutí",
                   [("Přezutí pneu", "Osobní i SUV"), ("Uskladnění", "Pneu i kola v suchu"),
                    ("Vyvážení kol", "Moderní technologie"), ("Prodej pneu", "Výhodné ceny")], "auto", "booking"),
    "autoškola": ("Řidičák snadno a v klidu", "Přihlásit se",
                  [("Skupina B", "Osobní automobil"), ("Kondiční jízdy", "Pro jistotu za volantem"),
                   ("Skupina A", "Motocykly"), ("Školení řidičů", "Profesní průkazy")], "auto", "booking"),
    "zubní ordinace": ("Zdravý a krásný úsměv", "Objednat se",
                       [("Preventivní prohlídky", "Pro děti i dospělé"), ("Dentální hygiena", "Šetrně a důkladně"),
                        ("Bělení zubů", "Zářivý úsměv"), ("Implantáty", "Moderní náhrady")], "medical", "booking"),
    "veterinář": ("Péče o vaše mazlíčky", "Objednat se",
                  [("Očkování", "Psi, kočky, králíci"), ("Preventivní prohlídky", "Zdravý mazlíček"),
                   ("Chirurgie", "Kastrace a zákroky"), ("Pohotovost", "Když spěcháte")], "medical", "booking"),
    "restaurace": ("Poctivá kuchyně a skvělá atmosféra", "Rezervovat stůl",
                   [("Polední menu", "Každý všední den"), ("Stálý lístek", "Sezónní suroviny"),
                    ("Rezervace stolu", "Online za pár vteřin"), ("Oslavy a akce", "Prostor pro vaše chvíle")],
                   "food", "menu"),
    "kavárna": ("Výběrová káva a domácí dorty", "Prohlédnout menu",
                [("Výběrová káva", "Z lokálních pražíren"), ("Domácí dorty", "Pečeme každý den"),
                 ("Snídaně", "Celý den"), ("Káva s sebou", "I do vlastního hrnku")], "food", "menu"),
    "pekárna": ("Čerstvé pečivo každé ráno", "Objednat online",
                [("Chléb", "Kváskový a poctivý"), ("Sladké pečivo", "Koláče, buchty, rohlíčky"),
                 ("Dorty na objednávku", "K oslavám i svatbám"), ("Catering", "Pro firmy i akce")], "food", "menu"),
    "květinářství": ("Květiny pro každou příležitost", "Objednat kytici",
                     [("Řezané květiny", "Čerstvé každý den"), ("Svatební kytice", "Na míru vašemu dni"),
                      ("Smuteční vazby", "Citlivě a důstojně"), ("Rozvoz květin", "Ještě dnes")], "flowers", "quote"),
    "fotograf": ("Fotky, které vyprávějí příběh", "Domluvit focení",
                 [("Svatby", "Celý den i obřad"), ("Portréty", "Osobní i firemní"),
                  ("Rodinné focení", "V ateliéru i venku"), ("Produktové foto", "Pro e-shopy")], "dark", "quote"),
    "instalatér": ("Voda, topení, plyn – rychle a spolehlivě", "Poptat opravu",
                   [("Havarijní služba", "Rychlý výjezd"), ("Koupelny", "Rekonstrukce na klíč"),
                    ("Topení", "Kotle a radiátory"), ("Odpady", "Čištění a opravy")], "crafts", "quote"),
    "elektrikář": ("Elektroinstalace bez starostí", "Poptat práci",
                   [("Elektroinstalace", "Byty, domy, firmy"), ("Revize", "Včetně zprávy"),
                    ("Opravy poruch", "Rychlý výjezd"), ("Chytrá domácnost", "Moderní řešení")], "crafts", "quote"),
    "truhlářství": ("Nábytek na míru z poctivého dřeva", "Nezávazná poptávka",
                    [("Kuchyně na míru", "Návrh i montáž"), ("Vestavěné skříně", "Každý centimetr využitý"),
                     ("Dveře a schody", "Masiv i dýha"), ("Renovace", "Nový život nábytku")], "crafts", "quote"),
    "malíř pokojů": ("Čisté malování bez stresu", "Nezávazná kalkulace",
                     [("Malování interiérů", "Rychle a čistě"), ("Stěrky", "Hladké stěny"),
                      ("Tapetování", "Moderní tapety"), ("Fasády", "Nátěry a opravy")], "crafts", "quote"),
    "stavební firma": ("Stavíme kvalitně a v termínu", "Poptat stavbu",
                       [("Rodinné domy", "Na klíč"), ("Rekonstrukce", "Byty i domy"),
                        ("Zateplení", "Úspora energie"), ("Projekce", "Od návrhu po kolaudaci")], "pro", "quote"),
    "účetní": ("Účetnictví, o které se nemusíte starat", "Nezávazná konzultace",
               [("Vedení účetnictví", "Pro OSVČ i firmy"), ("Daňová přiznání", "Včas a správně"),
                ("Mzdy", "Kompletní agenda"), ("Poradenství", "Daně bez stresu")], "pro", "quote"),
    "fitness": ("Tvoje forma začíná tady", "Vyzkoušet zdarma",
                [("Posilovna", "Moderní stroje"), ("Skupinové lekce", "Jóga, kruháč, spinning"),
                 ("Osobní trenér", "Plán na míru"), ("Permanentky", "Výhodné ceny")], "fitness", "booking"),
    "realitní kancelář": ("Prodáme vaši nemovitost rychle a výhodně", "Ocenit nemovitost",
                          [("Prodej", "Za nejlepší cenu"), ("Pronájem", "Spolehliví nájemníci"),
                           ("Odhad ceny", "Zdarma a do 24 h"), ("Právní servis", "Bezpečný převod")], "pro", "quote"),
    "úklidové služby": ("Čisto bez námahy", "Spočítat cenu",
                        [("Úklid domácností", "Pravidelně i jednorázově"), ("Úklid kanceláří", "Mimo pracovní dobu"),
                         ("Mytí oken", "Bez šmouh"), ("Po rekonstrukci", "Kompletní úklid")], "medical", "quote"),
    "zámečnictví": ("Zámečník, když ho potřebujete", "Zavolat ihned",
                    [("Otevírání dveří", "Nonstop výjezd"), ("Výměna zámků", "Bezpečnostní vložky"),
                     ("Kovové konstrukce", "Zábradlí, brány"), ("Bezpečnostní dveře", "Montáž a servis")],
                    "crafts", "quote"),
}
FALLBACK = ("Profesionální služby ve vašem okolí", "Kontaktovat",
            [("Služby", "Kvalitně a včas"), ("Poradenství", "Rádi poradíme"),
             ("Reference", "Spokojení zákazníci"), ("Kontakt", "Ozvěte se nám")], "pro", "quote")

def slugify(s):
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()
    s = re.sub(r"\b(s\.?r\.?o|spol|a\.?s|v\.?o\.?s|k\.?s)\b\.?", "", s)
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")[:40] or "firma"


def domain_for(lead):
    w = lead.get("website") or ""
    host = re.sub(r"^https?://", "", w).split("/")[0].lower()
    if host and not any(x in host for x in lf.SOCIAL + lf.FREE_BUILDERS):
        return host.replace("www.", "")
    return slugify(short_name(lead["name"])) + ".cz"


def short_name(name):
    n = re.sub(r",?\s*(s\.\s?r\.\s?o\.|spol\. s r\.\s?o\.|a\.\s?s\.|v\.\s?o\.\s?s\.)\s*$", "", name, flags=re.I)
    n = re.split(r"\s+[-–|]\s+", n)[0].strip()
    return n if len(n) <= 34 else n[:32].rsplit(" ", 1)[0] + "…"


def esc(s):
    return html.escape(str(s or ""))


# ---------------------------------------------------------------- лист
PROBLEM_CZ = {
    "БЕЗ САЙТУ": "Všiml jsem si ale, že nemáte vlastní webové stránky – zákazníci, kteří vás hledají na Googlu, "
                 "tak často skončí u konkurence.",
    "SOCIAL": "Všiml jsem si ale, že místo webu máte jen profil na sociální síti – na Googlu vás tak "
              "zákazníci najdou mnohem hůř.",
    "САЙТ НЕ ПРАЦЮЄ": "Zkoušel jsem otevřít váš web, ale nefunguje – zákazníci místo vaší nabídky vidí chybu.",
}
REASON_CZ = [
    ("телефон", "na mobilu se špatně zobrazuje (a většina zákazníků dnes hledá z telefonu)"),
    ("HTTPS", "prohlížeč u něj ukazuje „Nezabezpečeno“"),
    ("SSL", "prohlížeč hlásí problém s certifikátem"),
    ("©", "působí, že nebyl delší dobu aktualizován"),
    ("конструктор", "běží na bezplatném konstruktoru s cizí doménou"),
    ("CMS", "běží na zastaralém systému"),
    ("jQuery", "používá zastaralé technologie"),
    ("2000-х", "používá zastaralé technologie"),
    ("повільний", "načítá se pomalu"),
    ("порожній", "obsahuje velmi málo informací"),
]
PACKAGE_CZ = {
    "booking": "moderní web s online rezervací, ceníkem a galerií",
    "menu": "moderní web s menu, rezervací stolu a galerií",
    "quote": "moderní web s přehledem služeb, galerií prací a poptávkovým formulářem",
}


def email_text(lead, cfg):
    _, _, _, _, wkind = NICHES.get(lead.get("niche"), FALLBACK)
    name = short_name(lead["name"])
    status = lead["status"]
    reasons = lead.get("reasons") or ""
    if status == "БЕЗ САЙТУ":
        problem = PROBLEM_CZ["SOCIAL" if "лише" in reasons else "БЕЗ САЙТУ"]
    elif status == "САЙТ НЕ ПРАЦЮЄ":
        problem = PROBLEM_CZ[status]
    else:
        found = list(dict.fromkeys(cz for key, cz in REASON_CZ if key in reasons))[:2]
        problem = ("Prohlížel jsem si váš web a všiml jsem si, že " + " a ".join(found) + "."
                   if found else "Prohlížel jsem si váš web a myslím, že by si zasloužil modernější podobu.")
    reviews = int(lead.get("reviews") or 0)
    praise = f" Máte tam {reviews} hodnocení – to je skvělá vizitka." if reviews >= 10 else ""
    kc = lambda v: f"{int(v):,}".replace(",", " ")
    start, standard = kc(lead.get("price_start") or lf.START_PRICE), kc(lead.get("price") or lf.STANDARD_PRICE)
    support = int(lead.get("support") or 0)
    web = f" | {cfg['sender_web']}" if cfg.get("sender_web") else ""
    dom = domain_for(lead)
    own = lead.get("website") and dom in (lead.get("website") or "")
    domain_line = (f"web na vaší doméně {dom} se zabezpečením HTTPS" if own
                   else f"vlastní doménu {dom} a zabezpečení HTTPS")
    subject = f"Návrh nového webu pro {name}"
    body = f"""Dobrý den,

našel jsem {name} na Firmy.cz.{praise} {problem}

Připravil jsem pro vás ukázku, jak by mohl vypadat váš nový web – najdete ji v příloze.

Nabízím dva balíčky:
• Start – {start} Kč: jednostránkový web se službami, kontakty, mapou a formulářem
• Standard – {standard} Kč: {PACKAGE_CZ[wkind]} – přesně jako na ukázce

V obou balíčcích:
• perfektní zobrazení na mobilu
• {domain_line}
• základní nastavení pro Google, aby vás zákazníci našli

Platíte jednorázově, hotovo do {cfg['delivery_days']} dnů.
Správa a drobné úpravy potom za {support} Kč měsíčně (nepovinné).

Pokud vás to zaujme, stačí odpovědět na tento e-mail nebo zavolat na {cfg['sender_phone']}.

S pozdravem
{cfg['sender_name']}
{cfg['sender_phone']} | {cfg['sender_email']}{web}

P.S. Pokud o web nemáte zájem, stačí odepsat „ne“ a už se vám nebudu ozývat."""
    return subject, body, wa_text(lead, cfg)


# Конкретна вигода від сайту для кожної ніші (одне речення)
BENEFIT_CZ = {
    "booking": "Zákazníci se objednají sami online, i večer – vy nemusíte při práci zvedat telefon.",
    "autoservis": "Zákazníci si sami vyberou termín a uvidí ceník – méně telefonování, víc zakázek.",
    "pneuservis": "Zákazníci si sami zarezervují přezutí – žádné fronty a hovory v sezóně.",
    "autoškola": "Zájemci uvidí ceník a přihlásí se rovnou online – bez zbytečných dotazů.",
    "medical": "Klienti najdou ordinační hodiny, ceník a objednají se online – méně telefonátů.",
    "food": "Lidé uvidí menu a fotky a rovnou si rezervují stůl – víc hostů i ve všední dny.",
    "pekárna": "Lidé si objednají dorty a pečivo předem online – víc objednávek bez telefonování.",
    "květinářství": "Lidé si objednají kytici online i s rozvozem – i když máte zavřeno.",
    "crafts": "Kdo hledá řemeslníka, chce vidět fotky prací – web vám nosí poptávky, i když jste na zakázce.",
    "fotograf": "Klienti si vybírají fotografa podle portfolia – web s vašimi fotkami prodává za vás.",
    "účetní": "Firmy vybírají účetní podle důvěry – web s referencemi vám přivede nové klienty.",
    "realitní kancelář": "Prodávající hledají makléře online – vlastní web budí víc důvěry než inzerát.",
    "fitness": "Nováčci najdou rozvrh a ceník a přijdou rovnou na zkoušku.",
    "stavební firma": "Investoři chtějí vidět realizace – web s referencemi vám přivede větší zakázky.",
}
NICHE_GROUP = {
    "kadeřnictví": "booking", "kosmetický salon": "booking", "manikúra": "booking", "masáže": "booking",
    "barbershop": "booking", "zubní ordinace": "medical", "veterinář": "medical",
    "restaurace": "food", "kavárna": "food",
    "instalatér": "crafts", "elektrikář": "crafts", "truhlářství": "crafts", "malíř pokojů": "crafts",
    "zámečnictví": "crafts", "úklidové služby": "crafts",
}


def wa_text(lead, cfg):
    """Короткий WhatsApp: проблема → вигода → макет → м'яке питання."""
    name = short_name(lead["name"])
    niche = lead.get("niche") or ""
    place = (lead.get("city") or "").split(",")[-1].strip()
    search = f"„{niche} {place}“" if niche and place else "vaše služby"
    reasons = lead.get("reasons") or ""
    status = lead.get("status")
    if status == "БЕЗ САЙТУ" and "лише" in reasons:
        problem = f"všiml jsem si, že {name} má jen profil na sociální síti – kdo na Googlu hledá {search}, najde spíš konkurenci."
    elif status == "БЕЗ САЙТУ":
        problem = f"všiml jsem si, že {name} nemá web – kdo na Googlu hledá {search}, najde spíš konkurenci."
    elif status == "САЙТ НЕ ПРАЦЮЄ":
        problem = f"zkoušel jsem otevřít web {name} a nefunguje – zákazníci místo vás vidí chybu."
    elif "телефон" in reasons:
        problem = f"web {name} se na mobilu špatně zobrazuje, a přitom většina lidí hledá z telefonu."
    else:
        problem = f"web {name} působí zastarale a zákazníky to může odradit."
    benefit = BENEFIT_CZ.get(niche) or BENEFIT_CZ.get(NICHE_GROUP.get(niche, "crafts"))
    start = f"{int(lead.get('price_start') or lf.START_PRICE):,}".replace(",", " ")
    return (f"Dobrý den, tady {cfg['sender_name']}, dělám weby pro firmy v okolí. {problem[0].upper() + problem[1:]}\n\n"
            f"Připravil jsem vám ukázku, jak by váš web mohl vypadat (posílám obrázek). {benefit}\n\n"
            f"Web je hotový do týdne, cena od {start} Kč. Mám vám poslat podrobnosti?")


# ---------------------------------------------------------------- головне
def write_offers_xlsx(rows, path):
    cols = [("Пріоритет", "status", 15), ("Назва", "name", 32), ("Email", "email", 28),
            ("Телефон", "phone", 17), ("Start, Kč", "price_start", 9), ("Standard, Kč", "price", 10),
            ("Тема листа", "subject", 32),
            ("Текст листа", "body", 70), ("WhatsApp / SMS", "wa", 45), ("Картинка", "image_file", 30),
            ("Відправлено", "sent", 12)]
    keys = [k for _, k, _ in cols]
    wrap_cols = {keys.index("body") + 1, keys.index("wa") + 1}
    img_col = keys.index("image_file") + 1
    wb = Workbook()
    ws = wb.active
    ws.title = "Пропозиції"
    ws.append([c[0] for c in cols])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
    for r in rows:
        ws.append([r.get(k, "") for _, k, _ in cols])
        row = ws.max_row
        for col in range(1, len(cols) + 1):
            ws.cell(row, col).alignment = Alignment(vertical="top", wrap_text=col in wrap_cols)
        ws.cell(row, img_col).hyperlink = r["image_file"]
        ws.cell(row, img_col).font = Font(color="0563C1", underline="single")
        ws.row_dimensions[row].height = 120
    for i, (_, _, w) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def load_config():
    if not CONFIG.exists():
        CONFIG.write_text(json.dumps(DEFAULT_CONFIG, ensure_ascii=False, indent=2), "utf8")
        print(f"! Створив {CONFIG.name} – впиши туди своє ім'я, телефон, email.")
    return {**DEFAULT_CONFIG, **json.loads(CONFIG.read_text("utf8"))}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("leads_xlsx")
    ap.add_argument("--limit", type=int, default=0)
    ap.add_argument("--out", default=str(HERE / "offers"))
    ap.add_argument("--start", type=int, default=1, help="почати з ліда №")
    args = ap.parse_args()

    from playwright.sync_api import sync_playwright

    cfg = load_config()
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    leads = [lf.price(l) for l in load(args.leads_xlsx, "Ліди")]
    if args.limit:
        leads = leads[:args.limit]
    print(f"{len(leads)} лідів → {out}")

    import mockup
    html_dir = out / "_html"
    html_dir.mkdir(exist_ok=True)
    rows = []
    with sync_playwright() as p:
        exe = os.environ.get("CHROMIUM_PATH")
        browser = p.chromium.launch(executable_path=exe) if exe else p.chromium.launch()
        page = browser.new_page(viewport={"width": 1920, "height": 1080})
        for i, lead in enumerate(leads, 1):
            fname = out / f"{i:03d}_{slugify(lead['name'])}.jpg"
            if i >= args.start and not fname.exists():
                doc, combo = mockup.render(lead, domain_for(lead))
                hf = html_dir / (fname.stem + ".html")
                hf.write_text(doc, "utf8")
                page.goto(hf.as_uri(), wait_until="networkidle", timeout=90000)
                page.evaluate("document.fonts.ready")
                page.screenshot(path=str(fname), type="jpeg", quality=88)
            subject, body, wa = email_text(lead, cfg)
            rows.append({**lead, "subject": subject, "body": body, "wa": wa,
                         "image_file": str(fname.relative_to(out))})
            if i % 20 == 0:
                print(f"   {i}/{len(leads)}", flush=True)
        browser.close()

    write_offers_xlsx(rows, out / "nabidky.xlsx")
    print(f"Готово: {out / 'nabidky.xlsx'} + {len(rows)} картинок")


if __name__ == "__main__":
    main()
