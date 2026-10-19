"""
Пропозиції для британських лідів: унікальний макет + англійський лист (email) або сценарій дзвінка.

Запуск:  python uk_offers.py output/uk_leads.json
Результат: offers_uk/NNN.jpg + offers_uk/leads_uk.json (для пульта розсилки) + offers_uk/uk_offers.xlsx
"""
import json
import os
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import mockup
from offers import load_config, short_name, slugify
from uk_content import BENEFIT_EN, NICHES_EN, REASON_EN, START_GBP, STANDARD_GBP

HERE = Path(__file__).resolve().parent


def domain_of(lead):
    w = lead.get("website") or ""
    return re.sub(r"^https?://", "", w).split("/")[0].lower().replace("www.", "") or slugify(lead["name"]) + ".co.uk"


def problem_en(lead):
    reasons = " ".join(lead.get("reasons") or [])
    if lead["status"] == "САЙТ НЕ ПРАЦЮЄ":
        return f"I tried to visit {domain_of(lead)} and it isn’t loading – anyone searching for you right now hits an error page."
    found = list(dict.fromkeys(en for key, en in REASON_EN if key in reasons))[:2]
    if not found:
        return f"I had a look at {domain_of(lead)} and it feels a bit dated compared to what customers expect today."
    return f"I had a look at {domain_of(lead)} and noticed " + ", and ".join(found) + "."


def email_en(lead, cfg):
    name = short_name(lead["name"])
    label, tagline, cta, services, wkind = NICHES_EN[lead["niche"]]
    praise = (f" {lead['reviews']} reviews on Thomson Local is a great reputation to build on."
              if lead.get("reviews", 0) >= 5 else "")
    subject = f"A new website idea for {name}"
    body = f"""Hi there,

I came across {name} in {lead['city']}.{praise} {problem_en(lead)}

So I put together a quick mock-up of how a modern site for {name} could look – it’s attached.

{BENEFIT_EN[wkind]}

What you’d get:
• a fast, mobile-friendly site on your own domain ({domain_of(lead)})
• your services, reviews and photos front and centre
• a simple way for customers to call, book or request a quote
• set up for Google so local customers find you

Prices start at £{START_GBP} (one page) or £{STANDARD_GBP} for a full site like the mock-up, ready within 7 days. No monthly contracts.

Would it be worth a quick chat? Just reply to this email.

Best regards,
{cfg['sender_name']}
{cfg['sender_email']}

P.S. If it’s not for you, just reply “no” and I won’t contact you again."""
    reasons = " ".join(lead.get("reasons") or [])
    issue = ("it isn’t loading at the moment" if lead["status"] == "САЙТ НЕ ПРАЦЮЄ" else
             next((en for key, en in REASON_EN if key in reasons), "it looks a bit dated"))
    call = (f"Hi, is that {name}? My name is {cfg['sender_name']}, I build websites for local businesses. "
            f"I came across your website and noticed {issue}, so I’ve made a free mock-up of a new one for you. "
            f"Could I send it over by email or WhatsApp? A full site like that is £{STANDARD_GBP}, ready within a week.")
    return subject, body, call


def main():
    src = Path(sys.argv[1] if len(sys.argv) > 1 else "output/uk_leads.json")
    leads = json.loads(src.read_text("utf8"))
    out = HERE / "offers_uk"
    (out / "_html").mkdir(parents=True, exist_ok=True)
    cfg = load_config()
    from playwright.sync_api import sync_playwright

    rows = []
    with sync_playwright() as p:
        exe = os.environ.get("CHROMIUM_PATH")
        browser = p.chromium.launch(executable_path=exe) if exe else p.chromium.launch()
        page = browser.new_page(viewport={"width": 1920, "height": 1080})
        for i, lead in enumerate(leads, 1):
            lead.update(lang="en", rating=lead.get("stars") or "")
            img = out / f"{i:03d}.jpg"
            if not img.exists():
                doc, _ = mockup.render(lead, domain_of(lead))
                hf = out / "_html" / f"{i:03d}.html"
                hf.write_text(doc, "utf8")
                page.goto(hf.as_uri(), wait_until="networkidle", timeout=90000)
                page.evaluate("document.fonts.ready")
                page.screenshot(path=str(img), type="jpeg", quality=86)
            subject, body, call = email_en(lead, cfg)
            rows.append({
                "n": i, "name": short_name(lead["name"]), "full": lead["name"], "niche": NICHES_EN[lead["niche"]][0],
                "place": lead["city"], "status": {"САЙТ НЕ ПРАЦЮЄ": "DEAD SITE", "СТАРИЙ САЙТ": "OLD SITE"}.get(lead["status"], "WEAK SITE"),
                "phone": lead["phone"], "wa": lead["wa"], "phone_type": lead["phone_type"], "email": lead["email"],
                "channel": lead["channel"], "ltd": lead["ltd"], "reviews": lead["reviews"], "website": lead["website"],
                "subject": subject, "body": body, "watext": call, "slug": slugify(lead["name"]),
            })
            if i % 20 == 0:
                print(f"   {i}/{len(leads)}", flush=True)
        browser.close()

    (out / "leads_uk.json").write_text(json.dumps(rows, ensure_ascii=False), "utf8")
    wb = Workbook()
    ws = wb.active
    ws.title = "UK"
    cols = ["n", "name", "niche", "place", "status", "channel", "email", "phone", "website", "reviews", "subject", "body", "watext"]
    ws.append(cols)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="305496")
    for r in rows:
        ws.append([r[k] for k in cols])
        for c in ws[ws.max_row]:
            c.alignment = Alignment(vertical="top", wrap_text=c.column >= 11)
    wb.save(out / "uk_offers.xlsx")
    print(f"Готово: {len(rows)} пропозицій → {out}")


if __name__ == "__main__":
    main()
